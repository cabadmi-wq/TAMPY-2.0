# ==============================================
# TAMPY - Asistente de IA (Flask + Groq)
# ==============================================
import hmac
import io
import logging
import os
import re

from flask import Flask, jsonify, render_template, request
from groq import Groq

logging.basicConfig(level=logging.INFO)
log = logging.getLogger("tampy")

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 20 * 1024 * 1024  # 20 MB por petición

# ---------- Configuración (TODO viene de variables de entorno en Render) ----------
GROQ_API_KEY = os.environ.get("GROQ_API_KEY")
CLAVE_ACCESO = os.environ.get("TAMPY_CLAVE")  # opcional: protege tu app con una clave

client = Groq(api_key=GROQ_API_KEY) if GROQ_API_KEY else None

# Groq retira modelos seguido. TAMPY prueba en orden y salta al siguiente si uno ya no existe.
# Puedes forzar el tuyo con MODELO_TEXTO / MODELO_VISION en Render sin tocar el código.
def _lista(*items):
    return [m for m in items if m]

MODELOS_TEXTO = _lista(
    os.environ.get("MODELO_TEXTO"),
    "openai/gpt-oss-120b",
    "qwen/qwen3.8-27b",
    "openai/gpt-oss-20b",
)
MODELOS_VISION = _lista(
    os.environ.get("MODELO_VISION"),
    "qwen/qwen3.8-27b",
    "qwen/qwen3.6-27b",
    "meta-models/muse-spark-open",
)

SYSTEM = {
    "role": "system",
    "content": (
        "Eres TAMPY, un asistente de IA profesional, claro y amable. "
        "Respondes siempre en español, salvo que el usuario pida otro idioma. "
        "Usa formato Markdown (listas, negritas, bloques de código) cuando ayude. "
        "Si el usuario adjunta archivos o imágenes, básate en su contenido y no inventes "
        "lo que no aparezca en ellos."
    ),
}

MAX_MENSAJES = 12          # cuántos mensajes recientes se envían al modelo
MAX_CARACTERES = 60000     # tope total de texto enviado (cuida los límites de Groq)
MAX_TEXTO_ARCHIVO = 30000  # caracteres que se leen de cada archivo

EXT_TEXTO = {
    "txt", "md", "csv", "tsv", "json", "xml", "html", "htm", "css", "js", "ts", "py",
    "java", "c", "cpp", "cs", "php", "rb", "go", "rs", "sql", "yaml", "yml", "ini",
    "log", "sh", "bat",
}


# ---------- Utilidades ----------
def autorizado():
    if not CLAVE_ACCESO:
        return True
    return hmac.compare_digest(request.headers.get("X-Clave", ""), CLAVE_ACCESO)


def extraer_texto(nombre, datos):
    ext = nombre.lower().rsplit(".", 1)[-1] if "." in nombre else ""

    if ext == "pdf":
        from pypdf import PdfReader
        lector = PdfReader(io.BytesIO(datos))
        texto = "\n\n".join((p.extract_text() or "") for p in lector.pages[:80])
        if not texto.strip():
            raise ValueError("El PDF no tiene texto seleccionable (parece escaneado). "
                             "Súbelo como imagen para que TAMPY lo lea.")
    elif ext == "docx":
        import docx
        doc = docx.Document(io.BytesIO(datos))
        partes = [p.text for p in doc.paragraphs]
        for tabla in doc.tables:
            for fila in tabla.rows:
                partes.append(" | ".join(c.text.strip() for c in fila.cells))
        texto = "\n".join(partes)
    elif ext in ("xlsx", "xlsm"):
        from openpyxl import load_workbook
        libro = load_workbook(io.BytesIO(datos), read_only=True, data_only=True)
        partes = []
        for hoja in libro.worksheets[:6]:
            partes.append(f"## Hoja: {hoja.title}")
            for i, fila in enumerate(hoja.iter_rows(values_only=True)):
                if i >= 500:
                    partes.append("... (más filas omitidas)")
                    break
                partes.append(" | ".join("" if v is None else str(v) for v in fila))
        texto = "\n".join(partes)
    elif ext in EXT_TEXTO:
        for codec in ("utf-8-sig", "latin-1"):
            try:
                texto = datos.decode(codec)
                break
            except UnicodeDecodeError:
                continue
    else:
        raise ValueError(f"Formato .{ext or '?'} no soportado. Usa PDF, DOCX, XLSX, CSV, TXT, "
                         "código o imágenes.")

    truncado = len(texto) > MAX_TEXTO_ARCHIVO
    return texto[:MAX_TEXTO_ARCHIVO], truncado


def limpiar_mensajes(crudos):
    """Valida lo que manda el navegador y recorta el historial."""
    if not isinstance(crudos, list):
        raise ValueError("Formato de mensajes inválido.")
    limpios = []
    for m in crudos[-MAX_MENSAJES:]:
        rol = m.get("role")
        cont = m.get("content")
        if rol not in ("user", "assistant"):
            continue
        if isinstance(cont, str):
            limpios.append({"role": rol, "content": cont})
        elif isinstance(cont, list) and rol == "user":
            partes = []
            for p in cont:
                if p.get("type") == "text" and isinstance(p.get("text"), str):
                    partes.append({"type": "text", "text": p["text"]})
                elif p.get("type") == "image_url":
                    url = (p.get("image_url") or {}).get("url", "")
                    if url.startswith("data:image/"):
                        partes.append({"type": "image_url", "image_url": {"url": url}})
            if partes:
                limpios.append({"role": rol, "content": partes})

    # Imágenes solo en el último mensaje del usuario (ahorra tokens y evita errores)
    ultimo_user = max((i for i, m in enumerate(limpios) if m["role"] == "user"), default=-1)
    for i, m in enumerate(limpios):
        if isinstance(m["content"], list) and i != ultimo_user:
            texto = " ".join(p["text"] for p in m["content"] if p["type"] == "text")
            m["content"] = texto + "\n[Se adjuntó una imagen en este mensaje]"

    # Recorta por tamaño total, empezando por lo más viejo
    def tam(m):
        c = m["content"]
        return len(c) if isinstance(c, str) else sum(len(p.get("text", "")) for p in c)

    while len(limpios) > 1 and sum(tam(m) for m in limpios) > MAX_CARACTERES:
        limpios.pop(0)
    while limpios and limpios[0]["role"] != "user":
        limpios.pop(0)
    if not limpios:
        raise ValueError("No hay mensajes para procesar.")
    return limpios


def tiene_imagen(mensajes):
    return any(isinstance(m["content"], list) and
               any(p["type"] == "image_url" for p in m["content"]) for m in mensajes)


def modelo_no_disponible(error):
    t = str(error).lower()
    return any(s in t for s in ("model_not_found", "decommission", "does not exist",
                                "not found", "no longer supported", "does not support"))


def consultar(mensajes):
    candidatos = MODELOS_VISION if tiene_imagen(mensajes) else MODELOS_TEXTO
    ultimo_error = None
    for modelo in candidatos:
        try:
            r = client.chat.completions.create(
                model=modelo,
                messages=[SYSTEM] + mensajes,
                temperature=0.6,
                max_tokens=2048,
            )
            texto = (r.choices[0].message.content or "").strip()
            texto = re.sub(r"<think>.*?</think>", "", texto, flags=re.S).strip()
            return texto or "No obtuve respuesta del modelo.", modelo
        except Exception as e:  # noqa: BLE001
            ultimo_error = e
            if modelo_no_disponible(e):
                log.warning("Modelo %s no disponible: %s", modelo, e)
                continue
            raise
    raise RuntimeError(
        "Ningún modelo configurado está disponible en Groq. Revisa los nombres vigentes en "
        "console.groq.com/docs/models y ponlos en MODELO_TEXTO / MODELO_VISION. "
        f"Último error: {ultimo_error}"
    )


# ---------- Rutas ----------
@app.route("/")
def index():
    return render_template("index.html", requiere_clave=bool(CLAVE_ACCESO))


@app.route("/salud")
def salud():
    return "ok", 200


@app.route("/api/archivo", methods=["POST"])
def api_archivo():
    if not autorizado():
        return jsonify(error="Clave incorrecta."), 401
    f = request.files.get("archivo")
    if not f or not f.filename:
        return jsonify(error="No llegó ningún archivo."), 400
    try:
        texto, truncado = extraer_texto(f.filename, f.read())
    except ValueError as e:
        return jsonify(error=str(e)), 400
    except Exception as e:  # noqa: BLE001
        log.exception("Error leyendo archivo")
        return jsonify(error=f"No pude leer el archivo: {e}"), 400
    return jsonify(nombre=f.filename, texto=texto, truncado=truncado)


@app.route("/api/chat", methods=["POST"])
def api_chat():
    if not autorizado():
        return jsonify(error="Clave incorrecta."), 401
    if client is None:
        return jsonify(error="Falta la variable GROQ_API_KEY en el servidor."), 500
    try:
        mensajes = limpiar_mensajes((request.get_json(silent=True) or {}).get("mensajes"))
        texto, modelo = consultar(mensajes)
        return jsonify(respuesta=texto, modelo=modelo)
    except ValueError as e:
        return jsonify(error=str(e)), 400
    except Exception as e:  # noqa: BLE001
        log.exception("Error en /api/chat")
        t = str(e).lower()
        if "rate" in t and "limit" in t or "429" in t:
            return jsonify(error="Llegaste al límite de uso de Groq. Espera un minuto e intenta de nuevo."), 429
        if "invalid api key" in t or "401" in t:
            return jsonify(error="La API key de Groq es inválida o fue revocada."), 500
        return jsonify(error=f"Error del modelo: {e}"), 500


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 5000)), debug=False)
